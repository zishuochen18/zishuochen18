import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

wb = openpyxl.load_workbook('港澳流速/output/2026年6月港澳市场流速-初稿5.21.xlsx', data_only=True)
ws = wb['香港市场目标']

print("流速表汇总行数据:")
for row_idx in range(2, 20):
    b_val = ws.cell(row_idx, 2).value
    ap = ws.cell(row_idx, 42).value  # AP列
    au = ws.cell(row_idx, 47).value  # AU列
    if b_val in ("KOL-汇总", "商超&社群合计", "合计", "总计"):
        print(f"  {str(b_val):15} | 例子={ap} | 成本={au}")
