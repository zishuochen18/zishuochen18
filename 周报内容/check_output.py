import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

wb = openpyxl.load_workbook('港澳流速/output/2026年6月港澳市场流速-初稿5.21.xlsx')
ws = wb['香港市场目标']

print("输出流速表B列供应商列表(行2-15):")
for row_idx in range(2, 16):
    b_val = ws.cell(row_idx, 2).value
    ap_val = ws.cell(row_idx, 42).value  # AP列
    if b_val:
        print(f"  行{row_idx}: {b_val:20} | 例子={ap_val}")
