import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

wb = openpyxl.load_workbook('港澳流速/output/2026年6月港澳市场流速-初稿5.21.xlsx')
ws = wb['香港市场目标']

print("流速表B列所有数据行(行2-20):")
for row_idx in range(2, 21):
    b_val = ws.cell(row_idx, 2).value
    if b_val:
        ap = ws.cell(row_idx, 42).value
        print(f"  行{row_idx:2d}: {str(b_val):20} | AP={ap}")
