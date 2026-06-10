import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

wb = openpyxl.load_workbook('港澳流速/output/2026年6月港澳市场流速-初稿5.21.xlsx')
ws = wb['香港市场目标']

print("最终输出流速表B列供应商列表:")
for row_idx in range(2, 20):
    b_val = ws.cell(row_idx, 2).value
    ap_val = ws.cell(row_idx, 42).value  # AP列
    if b_val and str(b_val).strip() not in ("合计", ""):
        print(f"  行{row_idx}: {str(b_val):20} | 例子={ap_val}")
    if str(b_val).strip() == "合计":
        print(f"  行{row_idx}: {str(b_val):20} | 例子={ap_val}")
        break
