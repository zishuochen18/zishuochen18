import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

wb = openpyxl.load_workbook('港澳流速/output/2026年6月港澳市场流速-初稿5.21.xlsx')
ws = wb['香港市场目标']

print("流速表汇总行详细信息:")
for row_idx in range(2, 20):
    b_val = ws.cell(row_idx, 2).value
    if b_val in ("KOL-汇总", "商超&社群合计", "合计", "总计"):
        ap_cell = ws.cell(row_idx, 42)
        au_cell = ws.cell(row_idx, 47)
        print(f"\n行{row_idx}: {str(b_val):15}")
        print(f"  AP列(42): value={ap_cell.value}, formula={ap_cell.value if isinstance(ap_cell.value, str) else 'N/A'}")
        print(f"  AU列(47): value={au_cell.value}")
