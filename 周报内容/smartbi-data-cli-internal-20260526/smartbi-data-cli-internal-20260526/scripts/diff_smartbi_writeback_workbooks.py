#!/usr/bin/env python3
"""Build an offline diff manifest for SmartBI writeback workbooks."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from validate_smartbi_writeback_input import cell_text, get_task, load_config


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "configs/smartbi_writeback_tasks.json"
DEFAULT_KEY_COLUMNS = [
    "日期",
    "主投学科",
    "平台",
    "区域细分",
    "投放账户",
    "广告计划id",
    "广告组id",
    "广告id",
]


def row_key(row: dict[str, Any], key_columns: list[str]) -> str:
    return "|".join(cell_text(row.get(column)) for column in key_columns)


def load_rows(path: Path, task: dict[str, Any]) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    input_cfg = task.get("input") or {}
    schema = task.get("schema") or {}
    headers = [str(item) for item in schema.get("expected_headers", [])]
    workbook = load_workbook(path.expanduser(), read_only=False, data_only=True)
    sheet_name = str(input_cfg.get("sheet_name") or workbook.sheetnames[0])
    ws = workbook[sheet_name]
    header_row = int(input_cfg.get("header_row", 1))
    actual_headers = [cell_text(ws.cell(header_row, col).value) for col in range(1, len(headers) + 1)]
    if actual_headers != headers:
        raise ValueError(f"{path} header does not match configured template")
    data_start_row = int(input_cfg.get("data_start_row", header_row + 1))
    rows: list[tuple[int, dict[str, Any]]] = []
    for row_number in range(data_start_row, ws.max_row + 1):
        row = {header: ws.cell(row_number, col).value for col, header in enumerate(headers, start=1)}
        if all(cell_text(value) == "" for value in row.values()):
            continue
        rows.append((row_number, row))
    return headers, rows


def build_diff(
    *,
    original: Path,
    candidate: Path,
    task_name: str,
    task: dict[str, Any],
    max_changed_rows: int,
) -> dict[str, Any]:
    headers, original_rows = load_rows(original, task)
    _, candidate_rows = load_rows(candidate, task)
    schema = task.get("schema") or {}
    post_verify = task.get("post_verify") if isinstance(task.get("post_verify"), dict) else {}
    key_columns = [str(item) for item in post_verify.get("key_columns", []) if str(item) in headers]
    if not key_columns:
        key_columns = [str(item) for item in schema.get("key_columns", []) if str(item) in headers]
    if not key_columns:
        key_columns = DEFAULT_KEY_COLUMNS

    original_by_key: dict[str, tuple[int, dict[str, Any]]] = {}
    candidate_by_key: dict[str, tuple[int, dict[str, Any]]] = {}
    duplicate_original_keys = []
    duplicate_candidate_keys = []
    for row_number, row in original_rows:
        key = row_key(row, key_columns)
        if key in original_by_key:
            duplicate_original_keys.append(key)
        original_by_key[key] = (row_number, row)
    for row_number, row in candidate_rows:
        key = row_key(row, key_columns)
        if key in candidate_by_key:
            duplicate_candidate_keys.append(key)
        candidate_by_key[key] = (row_number, row)
    original_keys = set(original_by_key)
    candidate_keys = set(candidate_by_key)
    changed_rows = []
    blocked_columns = set(key_columns)
    blocked_columns.update(str(item) for item in schema.get("required_non_empty", []))
    allowed_value_columns = set(str(item) for item in schema.get("numeric_columns", []))
    errors: list[str] = []
    warnings: list[str] = []

    missing_keys = sorted(original_keys - candidate_keys)
    added_keys = sorted(candidate_keys - original_keys)
    if missing_keys:
        errors.append(f"candidate is missing {len(missing_keys)} original keyed rows")
    if added_keys:
        errors.append(f"candidate adds {len(added_keys)} keyed rows")
    if duplicate_original_keys:
        errors.append(f"original contains {len(duplicate_original_keys)} duplicate keys")
    if duplicate_candidate_keys:
        errors.append(f"candidate contains {len(duplicate_candidate_keys)} duplicate keys")

    for key in sorted(original_keys & candidate_keys):
        original_row_number, original_row = original_by_key[key]
        candidate_row_number, candidate_row = candidate_by_key[key]
        changes = []
        for column in headers:
            before = cell_text(original_row.get(column))
            after = cell_text(candidate_row.get(column))
            if before == after:
                continue
            if column in blocked_columns:
                errors.append(f"blocked column changed at candidate row {candidate_row_number}: {column}")
            if column not in allowed_value_columns:
                warnings.append(f"non-numeric column changed at candidate row {candidate_row_number}: {column}")
            changes.append({"column": column, "before": before, "after": after})
        if changes:
            changed_rows.append(
                {
                    "key": key,
                    "original_row": original_row_number,
                    "candidate_row": candidate_row_number,
                    "changes": changes,
                }
            )

    if len(changed_rows) > max_changed_rows:
        errors.append(f"changed rows {len(changed_rows)} exceeds max_changed_rows {max_changed_rows}")

    return {
        "status": "ok" if not errors else "error",
        "mode": "offline_writeback_diff",
        "boundary": {
            "no_smartbi_login": True,
            "no_upload": True,
            "no_external_write": True,
        },
        "task": task_name,
        "original": str(original),
        "candidate": str(candidate),
        "key_columns": key_columns,
        "changed_row_count": len(changed_rows),
        "max_changed_rows": max_changed_rows,
        "changed_rows": changed_rows,
        "errors": errors,
        "warnings": sorted(set(warnings)),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--task", required=True)
    parser.add_argument("--original", type=Path, required=True)
    parser.add_argument("--candidate", type=Path, required=True)
    parser.add_argument("--max-changed-rows", type=int, default=3)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--json", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    config = load_config(args.config.expanduser())
    task = get_task(config, args.task)
    result = build_diff(
        original=args.original.expanduser(),
        candidate=args.candidate.expanduser(),
        task_name=args.task,
        task=task,
        max_changed_rows=args.max_changed_rows,
    )
    if args.out:
        args.out.expanduser().parent.mkdir(parents=True, exist_ok=True)
        args.out.expanduser().write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        result["output"] = str(args.out)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"Status: {result['status']}")
        print(f"Changed rows: {result['changed_row_count']}")
    return 0 if result["status"] == "ok" else 2


if __name__ == "__main__":
    raise SystemExit(main())
