#!/usr/bin/env python3
"""Inspect SmartBI-exported workbooks with pandas-backed shape detection."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value)
    if text == "nan":
        return None
    return text.replace("\n", "").strip()


def is_present(value: Any) -> bool:
    cleaned = clean_cell(value)
    return cleaned not in (None, "")


def row_values(frame, index: int, max_columns: int = 30) -> list[Any]:
    values = [clean_cell(value) for value in frame.iloc[index].tolist()]
    return values[: min(len(values), max_columns)]


def non_empty_count(values: list[Any]) -> int:
    return sum(1 for value in values if value not in (None, ""))


def header_candidates(frame, max_rows: int = 20) -> list[dict[str, Any]]:
    candidates = []
    for index in range(min(len(frame), max_rows)):
        values = row_values(frame, index)
        count = non_empty_count(values)
        text_count = sum(1 for value in values if isinstance(value, str) and value)
        if count >= 2 and text_count >= 2:
            candidates.append({"row": index + 1, "non_empty_cells": count, "values": values})
    return candidates[:8]


def non_empty_ranges(frame) -> list[dict[str, int]]:
    row_numbers = []
    for index in range(len(frame)):
        if non_empty_count(row_values(frame, index, max_columns=len(frame.columns))) > 0:
            row_numbers.append(index + 1)
    ranges = []
    start = previous = None
    for row_number in row_numbers:
        if start is None:
            start = previous = row_number
        elif row_number == previous + 1:
            previous = row_number
        else:
            ranges.append({"start": start, "end": previous})
            start = previous = row_number
    if start is not None:
        ranges.append({"start": start, "end": previous})
    return ranges


def classify_sheet(frame, candidates: list[dict[str, Any]], merged_count: int) -> str:
    if not candidates:
        return "unknown"
    if merged_count == 0 and len(frame) > 20:
        column_count = max(len(frame.columns), 1)
        for candidate in candidates:
            if int(candidate["row"]) <= 6 and int(candidate["non_empty_cells"]) >= max(3, int(column_count * 0.6)):
                return "raw_table"
    first_header = int(candidates[0]["row"])
    rows_after = max(len(frame) - first_header, 0)
    if merged_count > 0 or len(candidates) >= 2 and int(candidates[1]["row"]) == first_header + 1:
        return "pivot_dashboard"
    if first_header <= 6 and rows_after >= 1:
        return "raw_table"
    return "unknown"


def openpyxl_metadata(path: Path) -> dict[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except Exception:
        return {}
    metadata = {}
    for ws in workbook.worksheets:
        metadata[ws.title] = {
            "openpyxl_max_row": ws.max_row,
            "openpyxl_max_column": ws.max_column,
            "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
        }
    return metadata


def inspect(path: Path) -> dict[str, Any]:
    try:
        import pandas as pd
    except ImportError as error:
        raise SystemExit("pandas is required; run with `uv run --with pandas --with openpyxl ...`.") from error

    metadata = openpyxl_metadata(path)
    workbook = pd.ExcelFile(path)
    result = {"path": str(path), "bytes": path.stat().st_size, "sheets": []}
    for sheet_name in workbook.sheet_names:
        frame = pd.read_excel(path, sheet_name=sheet_name, header=None, dtype=object)
        frame = frame.dropna(how="all").dropna(axis=1, how="all")
        candidates = header_candidates(frame)
        sheet_meta = metadata.get(sheet_name, {})
        merged_ranges = sheet_meta.get("merged_ranges", [])
        ranges = non_empty_ranges(frame)
        sample_rows = []
        for row_number in sorted({1, 2, 3, 4, 5, 6, len(frame)}):
            if 1 <= row_number <= len(frame):
                values = row_values(frame, row_number - 1)
                if non_empty_count(values):
                    sample_rows.append({"row": row_number, "values": values})
        result["sheets"].append(
            {
                "name": sheet_name,
                "pandas_rows": int(frame.shape[0]),
                "pandas_columns": int(frame.shape[1]),
                "openpyxl_max_row": sheet_meta.get("openpyxl_max_row"),
                "openpyxl_max_column": sheet_meta.get("openpyxl_max_column"),
                "merged_range_count": len(merged_ranges),
                "merged_ranges_sample": merged_ranges[:20],
                "header_candidates": candidates,
                "non_empty_ranges": ranges,
                "sample_rows": sample_rows,
                "likely_table_type": classify_sheet(frame, candidates, len(merged_ranges)),
            }
        )
    return result


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Inspect a SmartBI-exported workbook.")
    parser.add_argument("workbook")
    parser.add_argument("--out", help="Optional JSON output path.")
    parser.add_argument("--json", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    path = Path(args.workbook).expanduser()
    if not path.exists():
        raise SystemExit(f"Workbook not found: {path}")
    result = inspect(path)
    if args.out:
        out_path = Path(args.out).expanduser()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        result["output"] = str(out_path)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
