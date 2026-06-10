"""检查输出文件中公式是否保留"""
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

OUTPUT_DIR = Path(r"c:\Users\chenzishuo\Desktop\新建文件夹 (2)\周报内容\港澳流速\output")
OUTPUT_FILE = OUTPUT_DIR / "2026年5月港澳市场流速-初稿4.21.xlsx"

wb = load_workbook(OUTPUT_FILE)  # 不用 data_only，读取公式
ws = wb["香港市场目标"]

print("📌 检查公式保留情况（AO-AU列）")
col_names = ["AO", "AP", "AQ", "AR", "AS", "AT", "AU"]
for row in range(2, 11):
    supplier = ws.cell(row, 2).value
    if supplier is None:
        continue
    formulas = []
    for col_offset in range(7):
        col_idx = 41 + col_offset
        val = ws.cell(row, col_idx).value
        if isinstance(val, str) and val.startswith("="):
            formulas.append(f"{col_names[col_offset]}={val}")
    if formulas:
        print(f"  行{row:2} {str(supplier):15} 公式: {', '.join(formulas)}")
    else:
        print(f"  行{row:2} {str(supplier):15} 无公式（纯数值）")
