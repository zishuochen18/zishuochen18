"""检查行10-11的内容"""
import sys
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

SAMPLE_DIR = Path(r"c:\Users\chenzishuo\Desktop\新建文件夹 (2)\周报内容\港澳流速\sample")
FLOW_FILE = SAMPLE_DIR / "2026年5月港澳市场流速-初稿4.21.xlsx"

wb = load_workbook(FLOW_FILE)
ws = wb["香港市场目标"]

for row in range(6, 13):
    b_val = ws.cell(row, 2).value
    ao = ws.cell(row, 41).value
    ap = ws.cell(row, 42).value
    aq = ws.cell(row, 43).value
    ar = ws.cell(row, 44).value
    as_val = ws.cell(row, 45).value
    at = ws.cell(row, 46).value
    au = ws.cell(row, 47).value
    print(f"行{row:2}: B={str(b_val):15} AO={str(ao):20} AP={str(ap):15} AU={str(au):20}")
