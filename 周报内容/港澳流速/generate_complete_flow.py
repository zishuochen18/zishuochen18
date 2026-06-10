#!/usr/bin/env python3
"""
生成完整的6月港澳流速HTML报表
基于参考文件: 2026年6月港澳市场流速-初稿5.21.xlsx
"""
import sys
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).parent
SAMPLE_DIR = BASE_DIR / "sample"
OUTPUT_DIR = BASE_DIR / "output"

# 使用参考文件
FLOW_FILE = SAMPLE_DIR / "2026年6月港澳市场流速-初稿5.21.xlsx"


def extract_flow_data():
    """提取流速数据 - 从香港市场目标工作表"""
    wb = load_workbook(FLOW_FILE, data_only=True)
    ws = wb["香港市场目标"]

    rows_data = []
    # 从第2行开始读取（第1行是表头）
    for row_idx in range(2, 15):
        b_val = ws.cell(row_idx, 2).value  # B列：供应商名称
        if b_val is None:
            continue
        b_val = str(b_val).strip()

        # 读取各列数据
        # 假设列位置：AO=41(MTD目标), AP=42(实际), AR=44(约课目标), AS=45(约课实际), AU=47(成本)
        ao = ws.cell(row_idx, 41).value  # AO列：MTD例子目标
        ap = ws.cell(row_idx, 42).value  # AP列：实际例子达成
        ar = ws.cell(row_idx, 44).value  # AR列：MTD约课目标
        as_val = ws.cell(row_idx, 45).value  # AS列：实际约课达成
        au = ws.cell(row_idx, 47).value  # AU列：约课成本

        mtd_target = ao if isinstance(ao, (int, float)) else 0
        actual = ap if isinstance(ap, (int, float)) else 0
        lesson_actual = as_val if isinstance(as_val, (int, float)) else 0

        # 计算约课目标（如果没有明确值）
        if isinstance(ar, (int, float)):
            lesson_target = ar
        else:
            if b_val in ("代理商场", "书展"):
                lesson_target = mtd_target * 0.78
            elif b_val == "代理人汇总":
                lesson_target = mtd_target * 0.6
            elif b_val in ("KOL-汇总", "商超&社群合计", "合计"):
                lesson_target = 0
            else:
                lesson_target = mtd_target * 0.55

        rows_data.append({
            "name": b_val,
            "mtd_target": mtd_target,
            "actual": actual,
            "gap": actual - mtd_target,
            "lesson_target": round(lesson_target, 1),
            "lesson_actual": lesson_actual,
            "lesson_gap": round(lesson_actual - lesson_target, 1),
            "cost": au if isinstance(au, (int, float)) else None,
        })

    # 计算汇总行
    kol_names = ["钟嘉欣图片", "周家蔚", "钟嘉欣视频1", "其他汇总"]
    shangchao_names = ["代理商场", "书展", "代理人汇总"]

    kol_items = [r for r in rows_data if r["name"] in kol_names]
    shangchao_items = [r for r in rows_data if r["name"] in shangchao_names]

    for r in rows_data:
        if r["name"] == "KOL-汇总":
            r["mtd_target"] = sum(x["mtd_target"] for x in kol_items)
            r["actual"] = sum(x["actual"] for x in kol_items)
            r["gap"] = r["actual"] - r["mtd_target"]
            r["lesson_target"] = round(r["mtd_target"] * 0.55, 1)
            r["lesson_actual"] = sum(x["lesson_actual"] for x in kol_items)
            r["lesson_gap"] = round(r["lesson_actual"] - r["lesson_target"], 1)
        elif r["name"] == "商超&社群合计":
            r["mtd_target"] = sum(x["mtd_target"] for x in shangchao_items)
            r["actual"] = sum(x["actual"] for x in shangchao_items)
            r["gap"] = r["actual"] - r["mtd_target"]
            r["lesson_target"] = sum(x["lesson_target"] for x in shangchao_items)
            r["lesson_actual"] = sum(x["lesson_actual"] for x in shangchao_items)
            r["lesson_gap"] = round(r["lesson_actual"] - r["lesson_target"], 1)
        elif r["name"] == "合计":
            all_items = kol_items + shangchao_items
            r["mtd_target"] = sum(x["mtd_target"] for x in all_items)
            r["actual"] = sum(x["actual"] for x in all_items)
            r["gap"] = r["actual"] - r["mtd_target"]
            r["lesson_target"] = sum(x["lesson_target"] for x in all_items)
            r["lesson_actual"] = sum(x["lesson_actual"] for x in all_items)
            r["lesson_gap"] = round(r["lesson_actual"] - r["lesson_target"], 1)

    return rows_data


def generate_html(rows_data):
    """生成HTML报表"""
    html = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>6月港澳商务流速</title>
    <style>
        body {
            font-family: "Microsoft YaHei", Arial, sans-serif;
            margin: 20px;
            background-color: #f5f5f5;
        }
        .container {
            max-width: 1400px;
            margin: 0 auto;
            background-color: white;
            padding: 30px;
            border-radius: 8px;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        h1 {
            color: #333;
            border-bottom: 3px solid #4CAF50;
            padding-bottom: 10px;
        }
        .info {
            color: #666;
            margin-bottom: 20px;
            font-size: 14px;
        }
        table {
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
        }
        th, td {
            border: 1px solid #ddd;
            padding: 12px 8px;
            text-align: center;
        }
        th {
            background-color: #4CAF50;
            color: white;
            font-weight: bold;
        }
        tr:nth-child(even) {
            background-color: #f9f9f9;
        }
        tr:hover {
            background-color: #f1f1f1;
        }
        .supplier-name {
            text-align: left !important;
            font-weight: 500;
        }
        .num {
            text-align: right !important;
        }
        .positive {
            color: #4CAF50;
            font-weight: bold;
        }
        .negative {
            color: #f44336;
            font-weight: bold;
        }
        .summary-row {
            background-color: #e8f5e9 !important;
            font-weight: bold;
        }
        .total-row {
            background-color: #c8e6c9 !important;
            font-weight: bold;
            font-size: 1.05em;
        }
        .footer {
            margin-top: 30px;
            text-align: center;
            color: #999;
            font-size: 12px;
        }
    </style>
</head>
<body>
    <div class="container">
        <h1>📊 6月港澳商务流速</h1>
        <div class="info">
            <p><strong>生成时间：</strong>""" + datetime.now().strftime("%Y年%m月%d日 %H:%M:%S") + """</p>
            <p><strong>数据来源：</strong>2026年6月港澳市场流速-初稿5.21.xlsx</p>
        </div>

        <table>
            <thead>
                <tr>
                    <th>供应商</th>
                    <th>MTD例子目标</th>
                    <th>实际例子达成</th>
                    <th>例子Gap</th>
                    <th>MTD约课目标</th>
                    <th>实际约课达成</th>
                    <th>约课Gap</th>
                    <th>约课成本</th>
                </tr>
            </thead>
            <tbody>
"""

    for row in rows_data:
        # 判断是否为汇总行
        row_class = ""
        if row["name"] in ["KOL-汇总", "商超&社群合计"]:
            row_class = "summary-row"
        elif row["name"] == "合计":
            row_class = "total-row"

        # 判断gap的正负
        gap_class = "positive" if row["gap"] >= 0 else "negative"
        lesson_gap_class = "positive" if row["lesson_gap"] >= 0 else "negative"

        # 格式化成本
        cost_str = f"{row['cost']:.2f}" if row['cost'] is not None else "-"

        html += f"""                <tr class="{row_class}">
                    <td class="supplier-name">{row["name"]}</td>
                    <td class="num">{row["mtd_target"]:.0f}</td>
                    <td class="num">{row["actual"]:.0f}</td>
                    <td class="num {gap_class}">{row["gap"]:+.0f}</td>
                    <td class="num">{row["lesson_target"]:.1f}</td>
                    <td class="num">{row["lesson_actual"]:.0f}</td>
                    <td class="num {lesson_gap_class}">{row["lesson_gap"]:+.1f}</td>
                    <td class="num">{cost_str}</td>
                </tr>
"""

    html += """            </tbody>
        </table>

        <div class="footer">
            <p>自动生成 | 港澳商务流速数据</p>
        </div>
    </div>
</body>
</html>
"""

    return html


def main():
    print("=" * 60)
    print("6月港澳商务流速报表生成")
    print("=" * 60)
    print()

    if not FLOW_FILE.exists():
        print(f"[错误] 参考文件不存在: {FLOW_FILE}")
        return 1

    # 提取数据
    print("[步骤1] 提取流速数据...")
    rows_data = extract_flow_data()
    print(f"[信息] 提取了 {len(rows_data)} 行数据")

    # 生成HTML
    print("[步骤2] 生成HTML报表...")
    html = generate_html(rows_data)

    # 保存文件
    OUTPUT_DIR.mkdir(exist_ok=True)
    output_file = OUTPUT_DIR / f"6月港澳流速_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(html)

    print(f"\n[成功] HTML报表已生成: {output_file}")
    print(f"[成功] 请用浏览器打开查看")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    sys.exit(main())
