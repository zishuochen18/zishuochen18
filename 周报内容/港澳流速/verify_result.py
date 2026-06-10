"""验证处理结果与原表对比"""
import sys
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

SAMPLE_DIR = Path(r"c:\Users\chenzishuo\Desktop\新建文件夹 (2)\周报内容\港澳流速\sample")
OUTPUT_DIR = Path(r"c:\Users\chenzishuo\Desktop\新建文件夹 (2)\周报内容\港澳流速\output")
FLOW_FILE = SAMPLE_DIR / "2026年5月港澳市场流速-初稿4.21.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "2026年5月港澳市场流速-初稿4.21.xlsx"

# 对比原表和输出表的 AO-AU 列
wb_orig = load_workbook(FLOW_FILE, data_only=True)
wb_out = load_workbook(OUTPUT_FILE, data_only=True)
ws_orig = wb_orig["香港市场目标"]
ws_out = wb_out["香港市场目标"]

print("📊 原表 vs 输出表 对比（AO-AU列，行2-10）")
print(f"{'行':>3} {'供应商':15} {'列':8} {'原值':>10} {'新值':>10} {'匹配':>4}")
print("-" * 60)

col_names = ["AO:MTD目标", "AP:例子达成", "AQ:例子gap", "AR:MTD约课", "AS:约课达成", "AT:约课gap", "AU:约课成本"]

for row in range(2, 11):
    supplier = ws_orig.cell(row, 2).value
    if supplier is None:
        continue
    for col_offset, col_name in enumerate(col_names):
        col_idx = 41 + col_offset  # AO=41 in openpyxl 1-based
        orig_val = ws_orig.cell(row, col_idx).value
        new_val = ws_out.cell(row, col_idx).value

        # 格式化
        orig_str = f"{orig_val:.1f}" if isinstance(orig_val, float) else str(orig_val)
        new_str = f"{new_val:.1f}" if isinstance(new_val, float) else str(new_val)

        match = "✅" if orig_str == new_str else "⚠️"
        if orig_val is None and new_val is None:
            match = "✅"
        elif isinstance(orig_val, (int, float)) and isinstance(new_val, (int, float)):
            match = "✅" if abs((orig_val or 0) - (new_val or 0)) < 0.1 else "⚠️"

        if match == "⚠️":
            print(f"{row:>3} {str(supplier):15} {col_name:8} {orig_str:>10} {new_str:>10} {match}")

print("\n（只显示不匹配的行，全部匹配则无输出）")
