"""
从处理后的流速表中提取数据，生成 HTML 周报总结页面
"""
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

OUTPUT_DIR = Path(__file__).parent / "output"
FLOW_OUTPUT = OUTPUT_DIR / "2026年5月港澳市场流速-初稿4.21.xlsx"


def extract_data():
    """从处理后的流速表提取关键数据"""
    wb = load_workbook(FLOW_OUTPUT, data_only=True)
    ws = wb["香港市场目标"]

    rows_data = []
    for row_idx in range(2, 12):
        b_val = ws.cell(row_idx, 2).value
        if b_val is None:
            continue
        b_val = str(b_val).strip()

        ao = ws.cell(row_idx, 41).value  # MTD例子目标
        ap = ws.cell(row_idx, 42).value  # 实际例子达成
        ar = ws.cell(row_idx, 44).value  # MTD约课目标（公式，可能为None）
        as_val = ws.cell(row_idx, 45).value  # 实际约课达成
        au = ws.cell(row_idx, 47).value  # 约课成本

        mtd_target = ao if isinstance(ao, (int, float)) else 0
        actual = ap if isinstance(ap, (int, float)) else 0
        lesson_actual = as_val if isinstance(as_val, (int, float)) else 0

        # AR 列是公式（=AO*55% 或 =AO*78% 等），读不到值时自己算
        if isinstance(ar, (int, float)):
            lesson_target = ar
        else:
            # 根据供应商类型推算约课目标比例
            if b_val in ("代理商场", "书展"):
                lesson_target = mtd_target * 0.78
            elif b_val == "代理人汇总":
                lesson_target = mtd_target * 0.6
            elif b_val in ("KOL-汇总", "商超&社群合计", "合计"):
                lesson_target = 0  # 汇总行后面单独处理
            else:
                lesson_target = mtd_target * 0.55

        gap = actual - mtd_target
        lesson_gap = lesson_actual - lesson_target

        rows_data.append({
            "name": b_val,
            "mtd_target": mtd_target,
            "actual": actual,
            "gap": gap,
            "lesson_target": round(lesson_target, 1),
            "lesson_actual": lesson_actual,
            "lesson_gap": round(lesson_gap, 1),
            "cost": au if isinstance(au, (int, float)) else None,
        })

    # 补算汇总行的目标和 gap
    kol_names = ["钟嘉欣图片", "周家蔚", "钟嘉欣视频1", "其他汇总"]
    shangchao_names = ["代理商场", "书展", "代理人汇总"]

    kol_items = [r for r in rows_data if r["name"] in kol_names]
    shangchao_items = [r for r in rows_data if r["name"] in shangchao_names]

    for r in rows_data:
        if r["name"] == "KOL-汇总":
            r["mtd_target"] = sum(x["mtd_target"] for x in kol_items)
            r["actual"] = sum(x["actual"] for x in kol_items)
            r["gap"] = r["actual"] - r["mtd_target"]
            r["lesson_target"] = round(r["mtd_target"] * 0.55, 1)
            r["lesson_actual"] = sum(x["lesson_actual"] for x in kol_items)
            r["lesson_gap"] = round(r["lesson_actual"] - r["lesson_target"], 1)
        elif r["name"] == "商超&社群合计":
            r["mtd_target"] = sum(x["mtd_target"] for x in shangchao_items)
            r["actual"] = sum(x["actual"] for x in shangchao_items)
            r["gap"] = r["actual"] - r["mtd_target"]
            r["lesson_target"] = sum(x["lesson_target"] for x in shangchao_items)
            r["lesson_actual"] = sum(x["lesson_actual"] for x in shangchao_items)
            r["lesson_gap"] = round(r["lesson_actual"] - r["lesson_target"], 1)
        elif r["name"] == "合计":
            all_items = kol_items + shangchao_items
            r["mtd_target"] = sum(x["mtd_target"] for x in all_items)
            r["actual"] = sum(x["actual"] for x in all_items)
            r["gap"] = r["actual"] - r["mtd_target"]
            r["lesson_target"] = sum(x["lesson_target"] for x in all_items)
            r["lesson_actual"] = sum(x["lesson_actual"] for x in all_items)
            r["lesson_gap"] = round(r["lesson_actual"] - r["lesson_target"], 1)

    return rows_data


def calc_rate(actual, target):
    if target and target > 0:
        return round(actual / target * 100, 1)
    return 0


def generate_html(rows_data):
    """生成 HTML 周报总结"""
    # 分组
    kol_names = ["钟嘉欣图片", "周家蔚", "钟嘉欣视频1", "其他汇总"]
    kol_summary_name = "KOL-汇总"
    shangchao_names = ["代理商场", "书展", "代理人汇总"]
    shangchao_summary_name = "商超&社群合计"
    total_name = "合计"

    kol_rows = [r for r in rows_data if r["name"] in kol_names]
    kol_summary = next((r for r in rows_data if r["name"] == kol_summary_name), None)
    shangchao_rows = [r for r in rows_data if r["name"] in shangchao_names]
    shangchao_summary = next((r for r in rows_data if r["name"] == shangchao_summary_name), None)
    total_row = next((r for r in rows_data if r["name"] == total_name), None)

    # 找出落后和超预期的供应商
    behind = []
    ahead = []
    all_detail_rows = kol_rows + shangchao_rows
    for r in all_detail_rows:
        if r["gap"] < 0:
            behind.append(r)
        elif r["gap"] > 0:
            ahead.append(r)

    behind_lesson = [r for r in all_detail_rows if r["lesson_gap"] < 0]
    ahead_lesson = [r for r in all_detail_rows if r["lesson_gap"] > 0]

    today = datetime.now()
    yesterday = today - timedelta(days=1)
    date_str = yesterday.strftime("%Y年%m月%d日")

    def row_html(r, is_summary=False):
        rate = calc_rate(r["actual"], r["mtd_target"])
        lesson_rate = calc_rate(r["lesson_actual"], r["lesson_target"])
        gap_class = "negative" if r["gap"] < 0 else "positive" if r["gap"] > 0 else ""
        lgap_class = "negative" if r["lesson_gap"] < 0 else "positive" if r["lesson_gap"] > 0 else ""
        bold = "font-weight:700;" if is_summary else ""
        cost_str = f'{r["cost"]:.1f}' if r["cost"] else "-"
        return f"""<tr style="{bold}">
<td>{r["name"]}</td>
<td>{int(r["mtd_target"])}</td><td>{int(r["actual"])}</td>
<td class="{gap_class}">{int(r["gap"])}</td><td>{rate}%</td>
<td>{int(r["lesson_target"]) if r["lesson_target"] else '-'}</td><td>{int(r["lesson_actual"])}</td>
<td class="{lgap_class}">{int(r["lesson_gap"]) if r["lesson_gap"] else '-'}</td><td>{lesson_rate}%</td>
<td>{cost_str}</td>
</tr>"""

    # 生成 HTML
    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<title>港澳商务流速周报 - 截至{date_str}</title>
<style>
* {{ margin:0; padding:0; box-sizing:border-box; }}
body {{ font-family: -apple-system, "Microsoft YaHei", sans-serif; background:#f5f7fa; padding:24px; color:#333; }}
.container {{ max-width:1100px; margin:0 auto; }}
h1 {{ font-size:22px; margin-bottom:6px; }}
.subtitle {{ color:#666; font-size:14px; margin-bottom:24px; }}
.card {{ background:#fff; border-radius:10px; box-shadow:0 2px 8px rgba(0,0,0,0.06); padding:20px 24px; margin-bottom:20px; }}
.card h2 {{ font-size:16px; margin-bottom:12px; color:#1a1a1a; border-left:4px solid #4f8cff; padding-left:10px; }}
table {{ width:100%; border-collapse:collapse; font-size:13px; }}
th {{ background:#f0f4ff; padding:8px 6px; text-align:center; font-weight:600; border-bottom:2px solid #dde4f0; }}
td {{ padding:7px 6px; text-align:center; border-bottom:1px solid #eee; }}
tr:hover {{ background:#fafbff; }}
.negative {{ color:#e53935; font-weight:700; }}
.positive {{ color:#2e7d32; font-weight:700; }}
.alert-section {{ display:flex; gap:16px; flex-wrap:wrap; }}
.alert-card {{ flex:1; min-width:280px; padding:16px; border-radius:8px; }}
.alert-card.warn {{ background:#fff3e0; border:1px solid #ffe0b2; }}
.alert-card.good {{ background:#e8f5e9; border:1px solid #c8e6c9; }}
.alert-card h3 {{ font-size:14px; margin-bottom:8px; }}
.alert-card ul {{ list-style:none; padding:0; font-size:13px; }}
.alert-card li {{ padding:3px 0; }}
.tag {{ display:inline-block; padding:2px 8px; border-radius:4px; font-size:12px; margin-left:6px; }}
.tag-red {{ background:#ffcdd2; color:#c62828; }}
.tag-green {{ background:#c8e6c9; color:#2e7d32; }}
</style>
</head>
<body>
<div class="container">
<h1>港澳商务流速总结</h1>
<p class="subtitle">数据截至 {date_str}（MTD累计）</p>

<!-- 重点提醒 -->
<div class="card">
<h2>重点关注</h2>
<div class="alert-section">
<div class="alert-card warn">
<h3>未达进度（例子）</h3>
<ul>"""

    for r in sorted(behind, key=lambda x: x["gap"]):
        rate = calc_rate(r["actual"], r["mtd_target"])
        html += f'\n<li>{r["name"]} <span class="tag tag-red">gap {int(r["gap"])}</span> 达成率 {rate}%</li>'

    html += """
</ul>
</div>
<div class="alert-card warn">
<h3>未达进度（约课）</h3>
<ul>"""

    for r in sorted(behind_lesson, key=lambda x: x["lesson_gap"]):
        rate = calc_rate(r["lesson_actual"], r["lesson_target"])
        html += f'\n<li>{r["name"]} <span class="tag tag-red">gap {int(r["lesson_gap"])}</span> 达成率 {rate}%</li>'

    html += """
</ul>
</div>
<div class="alert-card good">
<h3>超预期达成</h3>
<ul>"""

    for r in sorted(ahead, key=lambda x: -x["gap"]):
        rate = calc_rate(r["actual"], r["mtd_target"])
        html += f'\n<li>{r["name"]} <span class="tag tag-green">+{int(r["gap"])}</span> 达成率 {rate}%</li>'

    html += """
</ul>
</div>
</div>
</div>

<!-- KOL 模块 -->
<div class="card">
<h2>KOL 模块</h2>
<table>
<tr><th>供应商</th><th>MTD目标</th><th>实际例子</th><th>例子gap</th><th>例子达成率</th><th>约课目标</th><th>实际约课</th><th>约课gap</th><th>约课达成率</th><th>约课成本</th></tr>
"""
    for r in kol_rows:
        html += row_html(r)
    if kol_summary:
        html += row_html(kol_summary, is_summary=True)

    html += """
</table>
</div>

<!-- 商超&社群 模块 -->
<div class="card">
<h2>商超&社群 模块</h2>
<table>
<tr><th>供应商</th><th>MTD目标</th><th>实际例子</th><th>例子gap</th><th>例子达成率</th><th>约课目标</th><th>实际约课</th><th>约课gap</th><th>约课达成率</th><th>约课成本</th></tr>
"""
    for r in shangchao_rows:
        html += row_html(r)
    if shangchao_summary:
        html += row_html(shangchao_summary, is_summary=True)

    html += """
</table>
</div>

<!-- 总计 -->
<div class="card">
<h2>总计</h2>
<table>
<tr><th>模块</th><th>MTD目标</th><th>实际例子</th><th>例子gap</th><th>例子达成率</th><th>约课目标</th><th>实际约课</th><th>约课gap</th><th>约课达成率</th><th>约课成本</th></tr>
"""
    if kol_summary:
        html += row_html(kol_summary)
    if shangchao_summary:
        html += row_html(shangchao_summary)
    if total_row:
        html += row_html(total_row, is_summary=True)

    html += """
</table>
</div>

</div>
</body>
</html>"""

    return html


def main():
    print("[1] 提取流速数据...")
    rows_data = extract_data()
    for r in rows_data:
        print(f"  {r['name']:15} 目标={r['mtd_target']}, 实际={r['actual']}, gap={r['gap']}")

    print(f"\n[2] 生成 HTML...")
    html = generate_html(rows_data)

    output_html = OUTPUT_DIR / "港澳商务流速总结.html"
    output_html.write_text(html, encoding="utf-8")
    print(f"\n[完成] 输出: {output_html}")


if __name__ == "__main__":
    main()
