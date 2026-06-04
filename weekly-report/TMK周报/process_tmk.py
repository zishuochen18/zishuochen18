"""
TMK 周报数据处理
1. 从 sample 目录读取三个 BI 导出的 Excel 文件
2. 解析合并数据，输出个人做工数据 Excel（带分组表头）
3. 提供 extract_tmk_data() 接口给 HTML 周报使用

参考自 zhoubao/generate_weekly_report.py
"""
"""
TMK 做工周报处理脚本（模块 3）

═══════════════════════════════════════════════════════════════
功能：合并 3 个 BI 报表数据，生成 TMK 个人做工汇总和异常列表
═══════════════════════════════════════════════════════════════

【输入】（3 个 BI 报表）
- sample/海外TMK做工监控.xlsx
- sample/海外TMK未邀约做工监控播报.xlsx
- sample/海外TMK做工勿扰情况汇总.xlsx

【输出】
- output/个人做工数据_{开始日期}_{结束日期}.xlsx（带分组表头 + 冻结窗格）

【数据合并逻辑】
按"TMK小组 + TMK"作为联合键合并 3 个表

【关键标准列定义】（STANDARD_COLUMNS）
- TMK 维度：日均通次、日均通次环比、日均通时、昨日通次、昨日通时
- 跟进时效：昨日生均跟进时效、生均跟进时效达成率
- 新生跟进：拨打学员数、生均通次、生均通时、有效接通率
- 老生跟进：同期拨打环比、生均通次、批量拨打次数、有效接通率
- 结果指标：思维例子数、思维约课数、首发本组约课数、思维例子约课率
- 新生跟进(进线非勿扰)：相关指标

【三个文件的表头位置】
- 做工监控：行 4 分组、行 5 列名、行 6 数据起
- 未邀约：行 6 分组、行 7 列名、行 8 数据起
- 勿扰：行 3 分组、行 4 列名、行 5 数据起

【异常判定规则】（ALERT_THRESHOLD = -0.10）
环比下降超过 10% 的指标会被标记为个人异常

【对外接口】
- extract_tmk_data() - 给 generate_weekly_report.py 使用
"""
import os
import sys
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).parent
SAMPLE_DIR = BASE_DIR / "sample"
OUTPUT_DIR = BASE_DIR / "output"

ALERT_THRESHOLD = -0.10  # 环比下降超 10% 视为异常

# 标准表格列定义：(分组名, 列名, 来源文件关键词, 数据格式)
STANDARD_COLUMNS = [
    ("TMK",       "日均通次",                   "做工监控", "int"),
    ("TMK",       "日均通次环比",               "做工监控", "pct"),
    ("TMK个人做工", "日均通时（分）",             "做工监控", "num"),
    ("TMK个人做工", "日均通时环比",               "做工监控", "pct"),
    ("TMK个人做工", "昨日通次",                   "做工监控", "int"),
    ("TMK个人做工", "昨日通时",                   "做工监控", "int"),
    ("跟进时效",   "昨日生均跟进时效（分）",       "做工监控", "num"),
    ("跟进时效",   "生均跟进时效（分）",           "做工监控", "num"),
    ("跟进时效",   "生均跟进时效达成率",           "做工监控", "pct"),
    ("新生跟进",   "拨打学员数",                  "做工监控", "int"),
    ("新生跟进",   "生均通次",                    "做工监控", "num"),
    ("新生跟进",   "生均通时（分）",              "做工监控", "num"),
    ("新生跟进",   "有效接通率",                  "做工监控", "pct"),
    ("老生跟进",   "同期拨打环比",                "做工监控", "pct"),
    ("老生跟进",   "生均通次",                    "做工监控", "num"),
    ("老生跟进",   "批量拨打次数",                "做工监控", "int"),
    ("老生跟进",   "有效接通率",                  "做工监控", "pct"),
    ("结果指标",   "思维例子数",                  "做工监控", "int"),
    ("结果指标",   "思维约课数",                  "做工监控", "int"),
    ("结果指标",   "首发本组约课数",              "做工监控", "int"),
    ("结果指标",   "思维例子约课率",              "做工监控", "pct"),
    ("新生跟进(进线非勿扰)", "生均通次",          "未邀约",   "num"),
    ("新生跟进(进线非勿扰)", "生均通次（进线非勿扰时段的例子）", "未邀约", "num"),
    ("新生跟进(进线非勿扰)", "生均通时（分）",     "未邀约",   "num"),
    ("新生跟进(进线非勿扰)", "首次接通邀约率",     "勿扰",     "pct"),
]


def find_excel(folder, keyword):
    for f in os.listdir(folder):
        if keyword in f and f.endswith(".xlsx") and not f.startswith("~$"):
            return os.path.join(folder, f)
    return ""


def parse_filter_params(df, max_rows=4):
    params = {}
    for i in range(min(max_rows, len(df))):
        row = df.iloc[i].tolist()
        for j, val in enumerate(row):
            if pd.notna(val):
                val_str = str(val).strip()
                if val_str in ("开始日期", "开始时间："):
                    if j + 1 < len(row) and pd.notna(row[j + 1]):
                        params["start_date"] = str(row[j + 1]).split(" ")[0]
                elif val_str in ("结束日期", "结束时间："):
                    if j + 1 < len(row) and pd.notna(row[j + 1]):
                        params["end_date"] = str(row[j + 1]).split(" ")[0]
                elif val_str == "一级渠道：":
                    if j + 1 < len(row) and pd.notna(row[j + 1]):
                        params["channel"] = str(row[j + 1])
                elif val_str == "二级渠道：":
                    if j + 1 < len(row) and pd.notna(row[j + 1]):
                        params["sub_channel"] = str(row[j + 1])
                elif val_str == "一级渠道":
                    if j + 1 < len(row) and pd.notna(row[j + 1]):
                        params.setdefault("channel", str(row[j + 1]))
                elif val_str == "二级渠道":
                    if j + 1 < len(row) and pd.notna(row[j + 1]):
                        params.setdefault("sub_channel", str(row[j + 1]))
    return params


def parse_excel_with_groups(filepath, header_group_row, header_name_row, data_start_row):
    raw = pd.read_excel(filepath, sheet_name="Sheet1", header=None)
    params = parse_filter_params(raw)

    group_row = raw.iloc[header_group_row].tolist()
    name_row = raw.iloc[header_name_row].tolist()

    column_meta = []
    current_group = ""
    for i in range(len(name_row)):
        if pd.notna(group_row[i]):
            current_group = str(group_row[i]).strip()
        if pd.notna(name_row[i]):
            column_meta.append((i, current_group, str(name_row[i]).strip()))

    metric_start = column_meta[0][0] if column_meta else 2

    columns = ["_skip"] * len(raw.columns)
    if metric_start >= 2:
        columns[metric_start - 2] = "TMK小组"
    columns[metric_start - 1] = "TMK"
    for idx, group, name in column_meta:
        if idx >= metric_start:
            full_name = f"{group}__{name}"
            columns[idx] = full_name

    seen = {}
    unique_columns = []
    for c in columns:
        if c in seen:
            seen[c] += 1
            unique_columns.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            unique_columns.append(c)

    data = raw.iloc[data_start_row:].copy()
    data.columns = unique_columns[:len(data.columns)]
    data = data.reset_index(drop=True)

    for col in data.columns:
        if col != "_skip" and col != "TMK小组" and col != "TMK":
            try:
                data[col] = pd.to_numeric(data[col], errors="coerce")
            except Exception:
                pass

    if "TMK" in data.columns:
        data = data[
            data["TMK"].notna()
            & ~data["TMK"].astype(str).str.match(r'^\d+）')
            & ~data["TMK"].astype(str).str.contains("口径|说明", na=False)
        ].reset_index(drop=True)

    if "TMK小组" in data.columns:
        data["TMK小组"] = data["TMK小组"].ffill()

    return params, data, column_meta


def parse_work_monitor(filepath):
    return parse_excel_with_groups(filepath, 4, 5, 6)


def parse_uninvited(filepath):
    return parse_excel_with_groups(filepath, 6, 7, 8)


def parse_disturb(filepath):
    return parse_excel_with_groups(filepath, 3, 4, 5)


def find_column_value(row, df, group, name):
    target = f"{group}__{name}"
    for col in df.columns:
        if col == target or col.startswith(target + "_"):
            return row.get(col)
    return None


def find_cross_file_row(target_df, group_name, person_name):
    if target_df is None or "TMK" not in target_df.columns:
        return None
    name_candidates = [person_name]
    if person_name == "团队汇总":
        name_candidates.append("总计")
    if "TMK小组" in target_df.columns and group_name:
        for n in name_candidates:
            row = target_df[
                (target_df["TMK"].astype(str).str.strip() == n)
                & (target_df["TMK小组"].astype(str).str.strip() == group_name)
            ]
            if len(row) > 0:
                return row.iloc[0]
    for n in name_candidates:
        row = target_df[target_df["TMK"].astype(str).str.strip() == n]
        if len(row) == 1:
            return row.iloc[0]
    return None


def infer_source_group(group, col_name, source):
    if source == "做工监控":
        if group == "TMK":
            return "TMK个人做工"
        return group
    elif source == "未邀约":
        if "新生跟进" in group:
            return "新生跟进"
        return group
    elif source == "勿扰":
        if any(k in col_name for k in ["约课", "邀约率", "首次接通", "首次拨打", "跟进时效", "例子数"]):
            return "结果指标"
        if "勿扰" in col_name:
            return "勿扰情况"
        return "结果指标"
    return group


def merge_person_data(work_df, uninvited_df, disturb_df):
    persons = work_df[work_df["TMK"].notna()].copy()
    rows = []
    for _, row in persons.iterrows():
        name = str(row["TMK"]).strip()
        group_name = str(row.get("TMK小组", "")).strip() if "TMK小组" in row else ""
        out = {"TMK小组": group_name, "TMK": name}
        for group, col_name, source, _ in STANDARD_COLUMNS:
            full = f"{group}__{col_name}"
            value = None
            if source == "做工监控":
                value = find_column_value(row, work_df, infer_source_group(group, col_name, "做工监控"), col_name)
            elif source == "未邀约":
                u_row = find_cross_file_row(uninvited_df, group_name, name)
                if u_row is not None:
                    src_group = infer_source_group(group, col_name, "未邀约")
                    value = find_column_value(u_row, uninvited_df, src_group, col_name)
            elif source == "勿扰":
                d_row = find_cross_file_row(disturb_df, group_name, name)
                if d_row is not None:
                    src_group = infer_source_group(group, col_name, "勿扰")
                    value = find_column_value(d_row, disturb_df, src_group, col_name)
            out[full] = value
        rows.append(out)
    return pd.DataFrame(rows)


def fmt(val, kind):
    if pd.isna(val) or val is None:
        return "-"
    try:
        v = float(val)
    except (ValueError, TypeError):
        return str(val)
    if kind == "pct":
        return f"{v * 100:.2f}%"
    elif kind == "int":
        return f"{int(round(v))}"
    else:
        return f"{v:.2f}"


def write_excel_cell(cell, value, kind):
    if value is None or pd.isna(value):
        cell.value = "-"
        return
    try:
        v = float(value)
    except (ValueError, TypeError):
        cell.value = str(value)
        return
    if kind == "pct":
        cell.value = v
        cell.number_format = "0.00%"
    elif kind == "int":
        cell.value = int(round(v))
        cell.number_format = "0"
    else:
        cell.value = round(v, 2)
        cell.number_format = "0.00"


def export_excel_table(work_df, uninvited_df, disturb_df, output_path):
    merged = merge_person_data(work_df, uninvited_df, disturb_df)

    groups_order = []
    group_to_cols = {}
    for group, col_name, _, kind in STANDARD_COLUMNS:
        if group not in groups_order:
            groups_order.append(group)
            group_to_cols[group] = []
        group_to_cols[group].append((col_name, kind))

    wb = Workbook()
    ws = wb.active
    ws.title = "个人做工数据"

    header_fill = PatternFill(start_color="DDEEFF", end_color="DDEEFF", fill_type="solid")
    sub_header_fill = PatternFill(start_color="EEF5FF", end_color="EEF5FF", fill_type="solid")
    total_fill = PatternFill(start_color="FFF4CC", end_color="FFF4CC", fill_type="solid")
    team_fill = PatternFill(start_color="E0E8F0", end_color="E0E8F0", fill_type="solid")
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.cell(row=1, column=1, value="TMK小组")
    ws.cell(row=1, column=2, value="TMK")
    ws.merge_cells(start_row=1, end_row=2, start_column=1, end_column=1)
    ws.merge_cells(start_row=1, end_row=2, start_column=2, end_column=2)

    col_idx = 3
    for g in groups_order:
        n = len(group_to_cols[g])
        ws.cell(row=1, column=col_idx, value=g)
        if n > 1:
            ws.merge_cells(start_row=1, end_row=1, start_column=col_idx, end_column=col_idx + n - 1)
        col_idx += n

    col_idx = 3
    for g in groups_order:
        for col_name, _ in group_to_cols[g]:
            ws.cell(row=2, column=col_idx, value=col_name)
            col_idx += 1

    total_columns = col_idx - 1

    for r in [1, 2]:
        for c in range(1, total_columns + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = header_fill if r == 1 else sub_header_fill
            cell.font = bold_font
            cell.alignment = center
            cell.border = border

    row = 3
    for _, person_row in merged.iterrows():
        ws.cell(row=row, column=1, value=person_row.get("TMK小组", "") or "")
        ws.cell(row=row, column=2, value=person_row.get("TMK", "") or "")
        col_idx = 3
        for g in groups_order:
            for col_name, kind in group_to_cols[g]:
                v = person_row.get(f"{g}__{col_name}")
                cell = ws.cell(row=row, column=col_idx)
                write_excel_cell(cell, v, kind)
                col_idx += 1

        tmk_val = str(person_row.get("TMK", "")).strip()
        group_val = str(person_row.get("TMK小组", "")).strip()
        is_grand_total = tmk_val == "总计" and group_val == "总计"
        is_team_total = tmk_val == "团队汇总"

        for c in range(1, total_columns + 1):
            ws.cell(row=row, column=c).border = border
            ws.cell(row=row, column=c).alignment = center
            if is_grand_total:
                ws.cell(row=row, column=c).fill = total_fill
                ws.cell(row=row, column=c).font = bold_font
            elif is_team_total:
                ws.cell(row=row, column=c).fill = team_fill
                ws.cell(row=row, column=c).font = bold_font
        row += 1

    for c in range(1, total_columns + 1):
        col_letter = get_column_letter(c)
        max_len = 4
        for r in range(1, row):
            cell = ws.cell(row=r, column=c)
            if cell.value:
                length = sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value))
                max_len = max(max_len, length)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 22)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "C3"

    wb.save(output_path)


def generate_overall_analysis(work_df):
    """从总计行生成整体数据要点"""
    points = []
    total_rows = work_df[work_df["TMK"].astype(str).str.contains("总计", na=False)]
    if len(total_rows) == 0:
        return points
    total = total_rows.iloc[0]

    metrics = {}
    for group, col_name, source, kind in STANDARD_COLUMNS:
        if source != "做工监控":
            continue
        v = find_column_value(total, work_df, infer_source_group(group, col_name, "做工监控"), col_name)
        if pd.notna(v):
            metrics[col_name] = float(v)

    return metrics


def analyze_alerts(work_df):
    """分析个人异常情况"""
    alerts = []
    persons = work_df[
        work_df["TMK"].notna()
        & ~work_df["TMK"].astype(str).str.contains("总计|团队汇总", na=False)
    ].copy()

    for _, row in persons.iterrows():
        name = str(row["TMK"]).strip()
        if not name or name == "nan":
            continue

        v = find_column_value(row, work_df, "TMK个人做工", "日均通次环比")
        if pd.notna(v) and float(v) < ALERT_THRESHOLD:
            alerts.append({
                "name": name,
                "type": "通次下降",
                "value": float(v),
                "desc": f"日均通次环比下降 {abs(float(v)) * 100:.0f}%"
            })

        v = find_column_value(row, work_df, "TMK个人做工", "日均通时环比")
        if pd.notna(v) and float(v) < ALERT_THRESHOLD:
            alerts.append({
                "name": name,
                "type": "通时下降",
                "value": float(v),
                "desc": f"日均通时环比下降 {abs(float(v)) * 100:.0f}%"
            })

        v = find_column_value(row, work_df, "跟进时效", "生均跟进时效达成率")
        if pd.notna(v) and float(v) < 0.3:
            alerts.append({
                "name": name,
                "type": "跟进时效低",
                "value": float(v),
                "desc": f"跟进时效达成率仅 {float(v) * 100:.0f}%"
            })

    return alerts


def extract_tmk_data():
    """供 HTML 周报使用的接口：返回完整的 TMK 周报数据"""
    work_file = find_excel(SAMPLE_DIR, "做工监控")
    uninvited_file = find_excel(SAMPLE_DIR, "未邀约")
    disturb_file = find_excel(SAMPLE_DIR, "勿扰")

    if not work_file:
        return None

    params, work_df, _ = parse_work_monitor(work_file)
    uninvited_df = None
    disturb_df = None
    if uninvited_file:
        _, uninvited_df, _ = parse_uninvited(uninvited_file)
    if disturb_file:
        _, disturb_df, _ = parse_disturb(disturb_file)

    merged = merge_person_data(work_df, uninvited_df, disturb_df)
    overall = generate_overall_analysis(work_df)
    alerts = analyze_alerts(work_df)

    # 收集分组结构
    groups_order = []
    group_to_cols = {}
    for group, col_name, _, kind in STANDARD_COLUMNS:
        if group not in groups_order:
            groups_order.append(group)
            group_to_cols[group] = []
        group_to_cols[group].append((col_name, kind))

    # 转换 merged 为字典列表
    rows = []
    for _, r in merged.iterrows():
        rows.append(r.to_dict())

    return {
        "params": params,
        "rows": rows,
        "groups_order": groups_order,
        "group_to_cols": group_to_cols,
        "overall": overall,
        "alerts": alerts,
        "standard_columns": STANDARD_COLUMNS,
    }


def main():
    SAMPLE_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)

    work_file = find_excel(SAMPLE_DIR, "做工监控")
    uninvited_file = find_excel(SAMPLE_DIR, "未邀约")
    disturb_file = find_excel(SAMPLE_DIR, "勿扰")

    if not work_file:
        print("[错误] 未找到 做工监控 Excel 文件")
        return

    print(f"[1] 解析做工监控: {os.path.basename(work_file)}")
    params, work_df, _ = parse_work_monitor(work_file)
    uninvited_df = None
    disturb_df = None
    if uninvited_file:
        print(f"[2] 解析未邀约: {os.path.basename(uninvited_file)}")
        _, uninvited_df, _ = parse_uninvited(uninvited_file)
    if disturb_file:
        print(f"[3] 解析勿扰: {os.path.basename(disturb_file)}")
        _, disturb_df, _ = parse_disturb(disturb_file)

    start_date = params.get("start_date", "?")
    end_date = params.get("end_date", "?")
    excel_name = f"个人做工数据_{start_date}_{end_date}.xlsx"
    excel_path = OUTPUT_DIR / excel_name

    try:
        export_excel_table(work_df, uninvited_df, disturb_df, str(excel_path))
        print(f"[完成] Excel 数据表: {excel_path}")
    except PermissionError:
        print(f"[错误] Excel 写入失败：{excel_path} 正被打开，请关闭后重试")
        return

    overall = generate_overall_analysis(work_df)
    alerts = analyze_alerts(work_df)

    print(f"\n[整体数据] 周期 {start_date} ~ {end_date}")
    if "日均通次" in overall:
        print(f"  团队日均通次: {int(overall['日均通次'])} 次")
    if "日均通次环比" in overall:
        v = overall["日均通次环比"]
        print(f"  环比: {v * 100:+.1f}%")
    if "日均通时（分）" in overall:
        print(f"  团队日均通时: {overall['日均通时（分）']:.1f} 分钟")
    if "生均跟进时效达成率" in overall:
        print(f"  生均跟进时效达成率: {overall['生均跟进时效达成率'] * 100:.1f}%")

    print(f"\n[个人异常] {len(alerts)} 项")
    for a in alerts:
        print(f"  - {a['name']}: {a['desc']}")


if __name__ == "__main__":
    main()
